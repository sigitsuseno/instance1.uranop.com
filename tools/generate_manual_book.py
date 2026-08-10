# -*- coding: utf-8 -*-
"""
Generator Manual Book dan Source Code — Aplikasi Uranop (Enterprise HRIS System).

Menghasilkan file .docx mengikuti struktur Manual Book PDKT (contoh).
Lokasi output: project root -> Manual_Book_dan_Source_Code_Uranop.docx

Cara pakai:
    python tools/generate_manual_book.py

Struktur dokumen:
    Cover -> Daftar Isi (field TOC otomatis) -> BAB I Pendahuluan
    -> BAB II Tools Program -> BAB III Manual Book & Source Code
    -> BAB IV Penutup.

Titik screenshot diberi kotak placeholder "[ SCREENSHOT DI SINI ]" —
pengguna tinggal mengganti dengan gambar screenshot sebenarnya.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------------------
# Konstanta warna & font
# ---------------------------------------------------------------------------
PRIMARY = RGBColor(0x1E, 0x40, 0xAF)      # biru Uranop
DARK    = RGBColor(0x1F, 0x29, 0x37)      # abu gelap
GREY    = RGBColor(0x64, 0x74, 0x8B)
CODE_BG = "EEF1F6"
PLACE_BG = "FFF7E6"
BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"
CODE_FONT = "Consolas"

doc = Document()

# ---------------------------------------------------------------------------
# Setup halaman A4 & margin
# ---------------------------------------------------------------------------
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

# ---------------------------------------------------------------------------
# Style dasar
# ---------------------------------------------------------------------------
def _set_font(style, name, size, bold=False, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    # pastikan font eastasia ikut ter-set (agar konsisten)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), name)
    rfonts.set(qn('w:hAnsi'), name)
    rfonts.set(qn('w:eastAsia'), name)

normal = doc.styles['Normal']
_set_font(normal, BODY_FONT, 11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color in [('Heading 1', 16, PRIMARY),
                          ('Heading 2', 13, DARK),
                          ('Heading 3', 11.5, GREY)]:
    st = doc.styles[name]
    _set_font(st, HEAD_FONT, size, bold=True, color=color)
    st.paragraph_format.space_before = Pt(14 if name != 'Heading 1' else 20)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

# ---------------------------------------------------------------------------
# Helper
# ---------------------------------------------------------------------------
def para(text, bold=False, italic=False, align=None, size=None, color=None,
         space_after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def _shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def _border(p, color="C9A24B", size="12"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '6')
        el.set(qn('w:color'), color)
        pbdr.append(el)
    pPr.append(pbdr)


def placeholder(keterangan):
    """Kotak placeholder untuk screenshot yang diisi user."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _shade(p, PLACE_BG)
    _border(p, color="C9A24B")
    r = p.add_run("\n[ SCREENSHOT DI SINI ]\n")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0xB0, 0x84, 0x00)
    r2 = p.add_run(keterangan)
    r2.italic = True
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0x8A, 0x6D, 0x1B)
    p.paragraph_format.space_after = Pt(10)


def code_block(code, lang_label=None):
    """Blok source code dengan latar abu-abu & font monospace."""
    if lang_label:
        cap = doc.add_paragraph()
        r = cap.add_run(lang_label)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        cap.paragraph_format.space_after = Pt(2)
        cap.paragraph_format.keep_with_next = True
    lines = code.strip('\n').split('\n')
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        _shade(p, CODE_BG)
        r = p.add_run(line if line else ' ')
        r.font.name = CODE_FONT
        r.font.size = Pt(7.5)
        rpr = r._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:ascii'), CODE_FONT)
        rfonts.set(qn('w:hAnsi'), CODE_FONT)
        rfonts.set(qn('w:cs'), CODE_FONT)
    # jarak setelah blok code
    doc.paragraphs[-1].paragraph_format.space_after = Pt(10)


def steps(items):
    """Langkah penggunaan — numbered list."""
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{i}. ")
        r.bold = True
        p.add_run(item)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_footer_pagenum():
    """Footer: nomor halaman di tengah."""
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld1 = OxmlElement('w:fldSimple')
    fld1.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = "1"
    run.append(t)
    fld1.append(run)
    p._p.append(fld1)


def add_toc():
    """Field TOC Word (isi otomatis setelah klik kanan -> Update Field)."""
    p = doc.add_paragraph()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = "Klik kanan di sini lalu pilih 'Update Field' untuk membuat daftar isi otomatis."
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = p.add_run()
    r._r.append(fldChar)
    r2 = p.add_run()
    r2._r.append(instrText)
    r3 = p.add_run()
    r3._r.append(fldChar2)
    r4 = p.add_run()
    r4._r.append(t)
    r5 = p.add_run()
    r5._r.append(fldChar3)


add_footer_pagenum()

# ===========================================================================
# COVER
# ===========================================================================
for _ in range(4):
    doc.add_paragraph()

para("MANUAL BOOK DAN SOURCE CODE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=24, color=PRIMARY)
para("APLIKASI URANOP", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=26, color=DARK)
para("Enterprise HRIS System", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=15, color=GREY, space_after=18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sistem Informasi Manajemen Sumber Daya Manusia Berbasis Web")
r.font.size = Pt(12)
r.font.color.rgb = GREY

for _ in range(6):
    doc.add_paragraph()

para("Mencakup Pengelolaan Data Karyawan, Kehadiran, Jadwal Kerja, Cuti, Kasbon,",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=GREY)
para("BPJS, PPh 21, Payroll, serta Laporan HR.", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=GREY)

for _ in range(4):
    doc.add_paragraph()

para("Dokumen ini berisi panduan penggunaan (manual book) dan potongan source code",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=GREY)
para("aplikasi Uranop untuk keperluan dokumentasi, pelatihan, dan pengembangan.",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=GREY)
para("Tahun 2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)

page_break()

# ===========================================================================
# DAFTAR ISI
# ===========================================================================
para("DAFTAR ISI", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, color=PRIMARY)
doc.add_paragraph()
add_toc()

page_break()

# ===========================================================================
# BAB I PENDAHULUAN
# ===========================================================================
doc.add_heading("BAB I PENDAHULUAN", level=1)

doc.add_heading("1. Latar Belakang", level=2)
para("Sumber daya manusia (SDM) merupakan aset terpenting bagi sebuah perusahaan atau organisasi. "
     "Pengelolaan data karyawan yang baik — mulai dari data pribadi, riwayat pekerjaan, kehadiran, "
     "cuti, kasbon, hingga perhitungan gaji — sangat menentukan kelancaran operasional dan kepatuhan "
     "terhadap peraturan ketenagakerjaan serta perpajakan.")
para("Sebelum adanya sistem terpadu, pengelolaan data HR pada umumnya masih dilakukan secara "
     "terpisah-pisah: data karyawan dicatat pada aplikasi spreadsheet, kehadiran diolah dari mesin "
     "fingerprint/face-scan secara manual, perhitungan lembur dan uang makan dihitung satu per satu, "
     "sedangkan perhitungan gaji, potongan BPJS, dan PPh 21 dilakukan dengan cara konvensional yang "
     "membutuhkan waktu lama dan rawan kesalahan. Data yang tersebar pada banyak berkas juga "
     "menyulitkan proses rekap, pelaporan, dan pengambilan keputusan manajemen.")
para("Uranop (Enterprise HRIS System) dibangun sebagai solusi terpadu untuk mengatasi permasalahan "
     "tersebut. Aplikasi berbasis web ini mengintegrasikan seluruh proses HR dalam satu sistem "
     "berbasis basis data (database), sehingga data dapat digunakan secara bersama oleh setiap bagian, "
     "konsisten, dan mudah ditelusuri. Aplikasi dilengkapi dengan otomatisasi perhitungan kehadiran, "
     "lembur, uang makan, kasbon, BPJS, PPh 21, hingga payroll dua tahap (on-the-fly dan on-record), "
     "serta beragam laporan yang dapat diekspor sesuai kebutuhan.")

doc.add_heading("2. Tujuan Pembuatan", level=2)
steps([
    "Mempermudah pengelolaan data karyawan secara terpusat dan terintegrasi dalam satu basis data.",
    "Mengotomatisasi perhitungan kehadiran, lembur, uang makan, kasbon, BPJS, PPh 21, dan payroll.",
    "Menyediakan alur persetujuan (approval) yang jelas untuk cuti, kasbon, dan lembur.",
    "Menghasilkan laporan HR yang cepat, akurat, dan dapat diekspor (Excel/PDF).",
    "Memberikan kontrol akses berbasis peran (role) bagi admin, HR, manajemen, dan supervisor.",
])

doc.add_heading("3. Ruang Lingkup", level=2)
para("Ruang lingkup aplikasi Uranop meliputi: Data Master (bagian, pekerjaan, kalender, konfigurasi "
     "gaji), Data Perusahaan, Pengelolaan Data Karyawan (CRUD, import, grouping, gaji, kontrak, "
     "kompensasi, keluarga, dokumen, resign/PHK), Jadwal Kerja (roster, pola kerja, shift), Kasbon, "
     "Cuti, Kehadiran (import, sync, lembur), BPJS, PPh 21, Payroll (gaji, slip, THR), serta "
     "berbagai laporan HR. Aplikasi memiliki dua area akses utama, yaitu area Admin dan area Supervisor.")

page_break()

# ===========================================================================
# BAB II TOOLS PROGRAM
# ===========================================================================
doc.add_heading("BAB II TOOLS PROGRAM", level=1)

tools = [
    ("1. Laravel (PHP Framework).",
     "Laravel adalah framework aplikasi web berbasis bahasa pemrograman PHP dengan sintaks yang "
     "ekspresif dan elegan. Laravel digunakan sebagai backend REST API Uranop, menangani autentikasi, "
     "validasi, logika bisnis, query basis data melalui Eloquent ORM, serta penyediaan endpoint "
     "JSON untuk aplikasi frontend. Arsitektur modul pada Uranop dikembangkan memanfaatkan struktur "
     "folder terpisah per modul (Auth, Employee, Payroll, dsb.) agar kode lebih terorganisasi."),
    ("2. Vue.js 3 (JavaScript Framework).",
     "Vue.js adalah framework JavaScript progresif untuk membangun antarmuka pengguna (user interface). "
     "Uranop menggunakan Vue 3 dengan Composition API (script setup) untuk membangun single-page "
     "application (SPA) yang interaktif dan responsif. Seluruh halaman di sisi frontend ditulis dalam "
     "berkas .vue yang memisahkan template, skrip, dan gaya."),
    ("3. Pinia (State Management).",
     "Pinia adalah pustaka manajemen state (state management) untuk Vue 3. Pinia digunakan untuk "
     "menyimpan data yang dipakai bersama antarmodul, misalnya data autentikasi (auth store), "
     "data permission, dan data yang perlu diakses lintas halaman tanpa pengambilan ulang dari server."),
    ("4. Vue Router.",
     "Vue Router adalah pustaka routing resmi untuk Vue.js. Pada Uranop, Vue Router mengatur navigasi "
     "antarhalaman dan melindungi akses berbasis peran (guard) — misalnya hanya pengguna dengan peran "
     "admin yang dapat mengakses area /admin, dan pengguna supervisor mengakses area /supervisor."),
    ("5. Tailwind CSS 4.",
     "Tailwind CSS adalah framework CSS utility-first untuk mempercepat pembuatan tampilan. Uranop "
     "menggunakan Tailwind untuk membangun desain antarmuka yang konsisten, modern, dan responsif "
     "dengan memanfaatkan variabel desain (CSS variable) untuk tema warna."),
    ("6. MySQL (Database).",
     "MySQL adalah sistem manajemen basis data relasional (RDBMS) yang digunakan untuk menyimpan "
     "seluruh data aplikasi. Struktur tabel dikelola melalui migration Laravel agar skema basis data "
     "terdokumentasi dan mudah dikembangkan."),
    ("7. Vite (Build Tool).",
     "Vite adalah build tool modern untuk proyek JavaScript yang cepat dalam mode pengembangan "
     "(hot module replacement). Vite dipakai untuk mengompilasi dan membundel aset Vue, CSS, "
     "dan JavaScript Uranop menjadi berkas statis yang siap produksi."),
    ("8. Pendukung Pengembangan.",
     "Composer (manajer dependensi PHP), Node.js & npm (untuk aset frontend), serta Laragon "
     "sebagai lingkungan pengembangan lokal di sistem operasi Windows. Seluruh kebutuhan aplikasi "
     "diinstal melalui composer.json dan package.json."),
]
for title, body in tools:
    doc.add_heading(title, level=2)
    para(body)

page_break()

# ===========================================================================
# BAB III MANUAL BOOK DAN SOURCE CODE
# ===========================================================================
doc.add_heading("BAB III MANUAL BOOK DAN SOURCE CODE", level=1)
para("Pada bab ini dijelaskan desain tampilan program, langkah-langkah penggunaan setiap fitur "
     "(manual book), beserta potongan source code dari modul terkait. Bagian yang ditandai dengan "
     "kotak [ SCREENSHOT DI SINI ] adalah tempat untuk meletakkan gambar screenshot halaman "
     "sebenarnya saat aplikasi sedang berjalan.")
doc.add_heading("1. Desain Tampilan Program", level=2)

# ---------------------------------------------------------------------------
# 1.1 Login
# ---------------------------------------------------------------------------
doc.add_heading("1.1. Tampilan Form Login", level=3)
placeholder("Tampilan halaman Login Uranop (input email, password, dan tombol Sign In).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Buka aplikasi Uranop pada peramban (browser) melalui alamat URL aplikasi.",
    "Masukkan alamat email pada kolom Email.",
    "Masukkan kata sandi pada kolom Password.",
    "Klik tombol Sign In. Jika email dan password benar, sistem akan mengarahkan ke halaman "
    "Dashboard sesuai peran pengguna (Admin atau Supervisor). Jika salah, akan muncul pesan "
    "\u201cEmail atau password tidak valid.\u201d",
])
doc.add_heading("Source Code — AuthApiController.php (Metode Login)", level=3)
code_block("""
namespace App\\Modules\\Auth\\Controllers\\Api\\V1;

use App\\Modules\\Auth\\Models\\User;
use App\\Modules\\Auth\\Resources\\AuthResource;
use Illuminate\\Http\\JsonResponse;
use Illuminate\\Http\\Request;
use Illuminate\\Routing\\Controller;
use Illuminate\\Support\\Facades\\Hash;
use Illuminate\\Validation\\ValidationException;

class AuthApiController extends Controller
{
    /** Web/Desktop login - hanya untuk HR/Admin. */
    public function login(Request $request): JsonResponse
    {
        $request->validate([
            'email'    => 'required|email',
            'password' => 'required',
        ]);

        $user = User::where('email', $request->email)->first();

        if (! $user || ! Hash::check($request->password, $user->password)) {
            throw ValidationException::withMessages([
                'email' => ['Email atau password tidak valid.'],
            ]);
        }

        if (! $user->is_active) {
            throw ValidationException::withMessages([
                'email' => ['Akun Anda dinonaktifkan. Hubungi administrator.'],
            ]);
        }

        // GATE: karyawan murni (tanpa role admin) ditolak
        $adminRoles = ['superadmin', 'hrmanager', 'adm_manager',
                       'hrbranch', 'hr_ast', 'manajemen'];
        $hasAdminRole = $user->hasAnyRole($adminRoles);
        if (! $hasAdminRole) {
            throw ValidationException::withMessages([
                'email' => ['Akses ditolak. Gunakan aplikasi mobile untuk login.'],
            ]);
        }

        $token = $user->createToken('api-token')->plainTextToken;
        $user->update([
            'last_login_at' => now(),
            'last_login_ip' => $request->ip(),
        ]);

        return response()->json([
            'token' => $token,
            'user'  => new AuthResource($user->load('roles', 'permissions')),
        ]);
    }
}
""", lang_label="Berkas: app/Modules/Auth/Controllers/Api/V1/AuthApiController.php")
code_block("""
<script setup>
import { ref } from 'vue'
import { useRouter } from 'vue-router'
import { useAuth } from '../../composables/useAuth'

const router = useRouter()
const auth = useAuth()

const form = ref({ email: '', password: '', remember: false })
const showPassword = ref(false)
const processing = ref(false)
const errors = ref({})

async function submit() {
  processing.value = true
  errors.value = {}
  try {
    await auth.login({ email: form.value.email, password: form.value.password })
    if (auth.canAccessAdmin) {
      router.push('/')
    } else {
      router.push('/supervisor')
    }
  } catch (e) {
    errors.value = { error: e.message || 'Login gagal. Periksa email dan password Anda.' }
  } finally {
    processing.value = false
  }
}
</script>
""", lang_label="Berkas: resources/js/Pages/Auth/Login.vue (bagian skrip)")

# ---------------------------------------------------------------------------
# 1.2 Dashboard
# ---------------------------------------------------------------------------
doc.add_heading("1.2. Tampilan Halaman Dashboard", level=3)
placeholder("Tampilan halaman Dashboard Admin/Supervisor (statistik ringkasan karyawan, kehadiran, cuti, dan lainnya).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Setelah berhasil login, sistem otomatis menampilkan halaman Dashboard.",
    "Dashboard menampilkan ringkasan data (jumlah karyawan aktif, kehadiran hari ini, "
    "pengajuan cuti/kasbon yang menunggu persetujuan, dan ringkasan lainnya).",
    "Gunakan menu pada sidebar kiri untuk berpindah ke modul lain.",
])

# ---------------------------------------------------------------------------
# 1.3 Data Master
# ---------------------------------------------------------------------------
doc.add_heading("1.3. Tampilan Data Master", level=3)
placeholder("Tampilan halaman Data Master (Bagian, Pekerjaan, Kalender, Gaji & LTHR).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Data Master pada sidebar, lalu pilih submenu (Bagian, Pekerjaan, Kalender, atau Gaji & LTHR).",
    "Untuk menambah data, klik tombol Tambah, isi formulir, lalu klik Simpan.",
    "Untuk mengubah data, klik ikon edit pada baris data, perbaiki isian, lalu klik Simpan.",
    "Untuk menghapus data, klik ikon hapus dan konfirmasi penghapusan.",
])

# ---------------------------------------------------------------------------
# 1.4 Data Karyawan
# ---------------------------------------------------------------------------
doc.add_heading("1.4. Tampilan Data Karyawan", level=3)
placeholder("Tampilan daftar Karyawan (tabel data karyawan dengan fitur pencarian dan filter).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Data Karyawan lalu submenu Karyawan untuk melihat daftar seluruh karyawan.",
    "Gunakan kotak pencarian dan filter (bagian, pekerjaan, status) untuk menyaring data.",
    "Untuk menambah karyawan baru, klik tombol Tambah, isi formulir (nama, NIK, jenis kelamin, "
    "bagian, jabatan, status kerja, tanggal masuk, dll.), lalu klik Simpan.",
    "Untuk melihat detail, klik baris karyawan atau tombol Lihat.",
    "Untuk mengubah data, klik tombol Edit, perbaiki data, lalu klik Simpan.",
    "Karyawan juga dapat diimpor massal melalui menu Import Karyawan dengan template Excel.",
    "Fitur Grouping Karyawan digunakan untuk mengelompokkan karyawan (misalnya berdasarkan grup gaji).",
])
doc.add_heading("Source Code — EmployeeApiController.php", level=3)
code_block("""
namespace App\\Modules\\Employee\\Controllers\\Api\\V1;

use App\\Http\\Controllers\\Controller;
use App\\Modules\\Employee\\Models\\Employee;
use App\\Modules\\Employee\\Resources\\EmployeeListResource;
use App\\Modules\\Employee\\Resources\\EmployeeResource;
use App\\Modules\\Employee\\Services\\EmployeeService;
use Illuminate\\Http\\JsonResponse;
use Illuminate\\Http\\Request;
use Illuminate\\Http\\Resources\\Json\\AnonymousResourceCollection;

class EmployeeApiController extends Controller
{
    public function __construct(
        protected EmployeeService $employeeService
    ) {}

    /** GET /api/employees - list karyawan dengan pagination dan filter. */
    public function index(Request $request): AnonymousResourceCollection
    {
        $filters = $request->validate([
            'search'            => 'nullable|string|max:100',
            'department_id'     => 'nullable|integer|exists:departments,id',
            'position_id'       => 'nullable|integer|exists:positions,id',
            'employment_status' => 'nullable|string',
            'is_active'         => 'nullable|boolean',
            'contract_type'     => 'nullable|string',
            'contract_status'   => 'nullable|string',
            'period_start'      => 'nullable|date',
            'period_end'        => 'nullable|date|after_or_equal:period_start',
            'per_page'          => 'nullable|integer|min:5|max:5000',
        ]);

        $employees = $this->employeeService->getPaginated($filters);

        return EmployeeListResource::collection($employees);
    }

    /** POST /api/employees */
    public function store(Request $request): JsonResponse
    {
        $data = $request->validate([
            'employee_code'          => 'nullable|string|max:50|unique:employees',
            'nik'                    => 'nullable|string|max:50|unique:employees',
            'name'                   => 'required|string|max:200',
            'gender'                 => 'required|in:L,P',
            'email'                  => 'nullable|email|max:200',
            'phone'                  => 'nullable|string|max:50',
            'employment_status'      => 'required|in:probation,contract,permanent,outsource,freelance',
            'join_date'              => 'required|date',
            'bank_name'              => 'nullable|string|max:100',
            'bank_account_number'    => 'nullable|string|max:100',
            'department_id'          => 'nullable|integer|exists:departments,id',
            'position_id'            => 'nullable|integer|exists:positions,id',
            'employee_group_codes'   => 'nullable|array',
            'employee_group_codes.*' => 'string',
        ]);

        $employee = $this->employeeService->create($data);

        return response()->json([
            'message' => "Karyawan {$employee->name} berhasil ditambahkan.",
            'data'    => new EmployeeResource(
                $employee->load(['department', 'position', 'latestContract', 'groups.master'])
            ),
        ], 201);
    }
}
""", lang_label="Berkas: app/Modules/Employee/Controllers/Api/V1/EmployeeApiController.php")

# ---------------------------------------------------------------------------
# 1.5 Jadwal Kerja
# ---------------------------------------------------------------------------
doc.add_heading("1.5. Tampilan Jadwal Kerja", level=3)
placeholder("Tampilan halaman Jadwal Kerja (Roster/Jadwal Umum, Pola & Jadwal Kerja, Shift).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Jadwal Kerja pada sidebar.",
    "Submenu Pola & Jadwal Kerja digunakan untuk menyusun pola kerja (work pattern) harian/mingguan.",
    "Submenu Shift digunakan untuk mengelola shift kerja (misalnya shift pagi, siang, malam).",
    "Submenu Jadwal Umum (Roster) menampilkan penempatan karyawan pada tanggal tertentu.",
    "Submenu Buat Jadwal digunakan untuk menghasilkan jadwal (roster) secara otomatis atau diimpor "
    "dari berkas template.",
])

# ---------------------------------------------------------------------------
# 1.6 Kasbon
# ---------------------------------------------------------------------------
doc.add_heading("1.6. Tampilan Pengelolaan Kasbon", level=3)
placeholder("Tampilan halaman Pengajuan Kasbon (pengajuan, persetujuan, pelunasan, riwayat).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Pengelolaan Kasbon lalu submenu Pengajuan Kasbon.",
    "Klik tombol Ajukan Kasbon, isi jumlah kasbon, alasan, dan jadwal pelunasan, lalu kirim pengajuan.",
    "Pengajuan disetujui/ditolak oleh atasan melalui submenu Persetujuan Kasbon.",
    "Pelunasan kasbon dilakukan secara otomatis sebagai potongan pada gaji melalui submenu Pelunasan Kasbon.",
    "Seluruh riwayat pengajuan dapat dilihat pada submenu Riwayat Kasbon.",
])

# ---------------------------------------------------------------------------
# 1.7 Cuti
# ---------------------------------------------------------------------------
doc.add_heading("1.7. Tampilan Pengelolaan Cuti", level=3)
placeholder("Tampilan halaman Pengajuan Cuti (generate cuti tahunan, pengajuan, pembatalan, saldo, rekap).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Pengelolaan Cuti.",
    "Generate Cuti Tahunan digunakan untuk membuat saldo cuti tahunan bagi seluruh karyawan.",
    "Pengajuan Cuti digunakan untuk mengajukan cuti (cuti tahunan, izin, sakit, dll.).",
    "Pembatalan Cuti digunakan untuk membatalkan cuti yang telah diajukan.",
    "Saldo Cuti menampilkan sisa jatah cuti setiap karyawan.",
    "Rekap Cuti menampilkan rekap pengambilan cuti dalam periode tertentu.",
])

# ---------------------------------------------------------------------------
# 1.8 Kehadiran
# ---------------------------------------------------------------------------
doc.add_heading("1.8. Tampilan Kehadiran", level=3)
placeholder("Tampilan halaman Kehadiran (import kehadiran, manual sync, sync kehadiran, hitung lembur, resume).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Kehadiran pada sidebar.",
    "Import Kehadiran digunakan untuk mengunggah data log kehadiran (dari mesin absensi).",
    "Manual Sync dan Sync Kehadiran digunakan untuk menyinkronkan data kehadiran secara manual/otomatis.",
    "Hitung Lembur digunakan untuk memproses perhitungan lembur dari data kehadiran.",
    "Resume Kehadiran menampilkan ringkasan kehadiran karyawan pada periode tertentu.",
    "Consecutive Day menampilkan informasi hari kerja beruntun (berguna untuk deteksi kelelahan/aturan).",
])

# ---------------------------------------------------------------------------
# 1.9 BPJS
# ---------------------------------------------------------------------------
doc.add_heading("1.9. Tampilan Pengelolaan BPJS", level=3)
placeholder("Tampilan halaman BPJS (Keanggotaan, Iuran BPJS, Konfigurasi BPJS).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Pengelolaan BPJS.",
    "Submenu Keanggotaan menampilkan data keanggotaan BPJS Ketenagakerjaan dan Kesehatan karyawan.",
    "Submenu Iuran BPJS menampilkan perhitungan iuran BPJS (JHT, JP, JKK, JKM, JKN) per karyawan.",
    "Submenu Konfigurasi BPJS digunakan untuk mengatur persentase iuran BPJS.",
])

# ---------------------------------------------------------------------------
# 1.10 PPh 21
# ---------------------------------------------------------------------------
doc.add_heading("1.10. Tampilan PPh 21", level=3)
placeholder("Tampilan halaman PPh 21 (Pajak Karyawan, TER Bulanan, PPh 21 Tahunan).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu PPh 21 pada sidebar.",
    "Pajak Karyawan menampilkan data status pajak (PTKP, NPWP) setiap karyawan.",
    "TER Bulanan digunakan untuk mengelola tarif efektif rata-rata (TER) pajak bulanan.",
    "PPh 21 Tahunan digunakan untuk menghitung perhitungan PPh 21 tahunan.",
])

# ---------------------------------------------------------------------------
# 1.11 Payroll
# ---------------------------------------------------------------------------
doc.add_heading("1.11. Tampilan Payroll", level=3)
placeholder("Tampilan halaman Payroll — Gaji Karyawan (perhitungan gaji dua mode: on-the-fly dan on-record).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Payroll lalu submenu Gaji Karyawan.",
    "Pilih periode penggajian. Jika periode belum berakhir, sistem menghitung gaji secara "
    "on-the-fly (langsung dari data kehadiran). Jika sudah berakhir, sistem membaca hasil simpanan "
    "(pay_record).",
    "Klik tombol Simpan untuk membuat snapshot hasil perhitungan (status draft).",
    "Periksa dan perbaiki data (misalnya jumlah lembur, kasbon, atau cabang bank).",
    "Setelah periode berakhir, klik Finalisasi untuk menghitung ulang dengan rumus final (status generated).",
    "Klik Kunci (Lock) agar data payroll tidak dapat diubah lagi. Buka kunci (Unlock) hanya "
    "dengan password yang telah diatur.",
    "Slip Gaji menampilkan rincian slip gaji karyawan; Perhitungan THR digunakan untuk menghitung "
    "Tunjangan Hari Raya.",
])
doc.add_heading("Source Code — GajiKaryawanController.php", level=3)
code_block("""
class GajiKaryawanController extends Controller
{
    // Lifecycle status: draft -> generated -> locked
    private const STATUS_DRAFT     = 'draft';
    private const STATUS_GENERATED = 'generated';
    private const STATUS_LOCKED    = 'locked';

    /** GET /api/v1/payroll/gaji-karyawan?period_id=X&segment=Y */
    public function index(Request $request)
    {
        $validated = $request->validate([
            'period_id' => 'required|exists:pay_periods,id',
            'segment'   => 'nullable|in:A,B',
        ]);

        $period  = PayPeriod::findOrFail($validated['period_id']);
        $segment = $this->resolveSegment($period, $validated['segment'] ?? null);

        // End_date belum lewat  -> hitung live (on-the-fly)
        // End_date sudah lewat  -> baca pay_records (on-record)
        if ($period->end_date && Carbon::today()->lte($period->end_date)) {
            return $this->onTheFly($period, $segment);
        }
        return $this->onRecord($period, $segment);
    }

    /** Hitung gaji per karyawan (potongan inti rumus on-the-fly). */
    private function computeOnTheFlyRows(PayPeriod $period, ?string $segment): array
    {
        $fixedDays = $this->fixedWorkingDay();           // pembagi (default 25)
        $holidays  = Holiday::whereBetween('date', [
            $period->start_date->toDateString(),
            $period->end_date->toDateString(),
        ])->pluck('date')->map(fn ($d) => $d->toDateString())->toArray();

        $employees = Employee::with(['department', 'position', 'groups', 'bpjs'])
            ->activeInPeriod(
                $period->start_date->toDateString(),
                $period->end_date->toDateString()
            )
            ->whereHas('shiftRosters', fn ($q) => $q->whereBetween('date', [
                $period->start_date->toDateString(),
                $period->end_date->toDateString(),
            ]))
            ->get()
            ->filter(fn ($emp) => $emp->isGroupGaji())
            ->sortBy('no_urut')->sortBy('nip')->values();

        foreach ($employees as $employee) {
            $prepares = AttendancePrepare::where('employee_id', $employee->id)
                ->whereBetween('date', [$segStart, $effEnd])->get();

            $hariKerja = $prepares->filter(function ($p) use ($holidays) {
                $date = $p->date->toDateString();
                if (Carbon::parse($date)->isSunday()) return false;   // Minggu
                if (in_array($date, $holidays)) return false;          // Libur
                if ($p->status === AttendancePrepare::STATUS_ABSENT) return false;
                return true;
            })->count();

            $gajiPokok   = $employee->gaji_pokok($segmentMonth);
            $premi       = $employee->premi($segmentMonth);
            $tjMasaKerja = $employee->tunjangan_masa_kerja($segmentMonth);
            $tunjangan   = $employee->tunjangan($segmentMonth);

            $gaji = round(($gajiPokok / $fixedDays) * $hariKerja, 2);
            $premiHadir = round(($premi / $fixedDays) * $hariKerja, 2);

            // Upah lembur: basis per-jam /173, dibulatkan ke kelipatan 100
            $upahLembur = $totalLemburJam > 0
                ? ceil((($gajiPokok + $tjMasaKerja + $tunjangan) / 173)
                        * $totalLemburJam / 100) * 100
                : 0;

            $gajiKotor = $gaji + $tunjangan + $upahLembur + $premiHadir + $revisi;
            $beforeRounding = $gajiKotor - ($bpjsTk + $bpjsKs + $bpjsPen + $pph + $cashbon);
            $rounded  = ceil($beforeRounding / 100) * 100;
            $pblt     = round($rounded - $beforeRounding, 2);
            $gajiBersih = $rounded;
        }
        return ['rows' => $rows, 'fixed_days' => $fixedDays];
    }
}
""", lang_label="Berkas: app/Modules/Payroll/Controllers/Api/V1/GajiKaryawanController.php (potongan)")
code_block("""
Route::prefix('v1/payroll')->middleware(['api'])->group(function () {
    Route::apiResource('periods', PayPeriodApiController::class);
    Route::get('gaji-karyawan', [GajiKaryawanController::class, 'index']);

    // Lifecycle payroll 2-mode: Simpan (snapshot) -> Finalisasi -> Lock -> Unlock
    Route::post('gaji-karyawan/simpan',     [GajiKaryawanController::class, 'simpan']);
    Route::post('gaji-karyawan/finalisasi', [GajiKaryawanController::class, 'finalisasi']);
    Route::post('gaji-karyawan/lock',       [GajiKaryawanController::class, 'lock']);
    Route::post('gaji-karyawan/unlock',     [GajiKaryawanController::class, 'unlock']);

    Route::put('gaji-karyawan/{id}/upah-lembur', [GajiKaryawanController::class, 'updateUpahLembur']);
    Route::put('gaji-karyawan/{id}/transfer-info', [GajiKaryawanController::class, 'updateTransferInfo']);
    Route::put('gaji-karyawan/bulk-update-cabang', [GajiKaryawanController::class, 'bulkUpdateCabang']);

    Route::get('configs/{type}', [PayrollConfigApiController::class, 'show']);
    Route::put('configs/{type}', [PayrollConfigApiController::class, 'update']);
    Route::get('payslips',       [PayslipController::class, 'index']);
    Route::get('thr',            [ThrApiController::class, 'index']);
    Route::post('thr/generate',  [ThrApiController::class, 'generate']);
});
""", lang_label="Berkas: app/Modules/Payroll/Routes/api.php")

# ---------------------------------------------------------------------------
# 1.12 Laporan
# ---------------------------------------------------------------------------
doc.add_heading("1.12. Tampilan Laporan", level=3)
placeholder("Tampilan halaman Laporan (Lembur & Uang Makan, Kehadiran, Payroll, Pajak, BPJS, Rekap).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Laporan pada sidebar.",
    "Pilih jenis laporan yang diinginkan (Lembur & Uang Makan, Kehadiran, Payroll, Pajak, BPJS, "
    "Rekap Uang Makan, Rekap Gaji, Rekap Kerja, Rekap PPh & Kompensasi, Karyawan Titipan).",
    "Atur rentang periode/filter pada halaman laporan.",
    "Klik tombol Ekspor untuk mengunduh laporan dalam format Excel atau format lain yang tersedia.",
])

# ---------------------------------------------------------------------------
# 1.13 Settings
# ---------------------------------------------------------------------------
doc.add_heading("1.13. Tampilan Pengaturan (Settings)", level=3)
placeholder("Tampilan halaman Pengaturan aplikasi (konfigurasi umum, notifikasi, keamanan).")
doc.add_heading("Langkah penggunaan :", level=3)
steps([
    "Pilih menu Settings pada sidebar (hanya untuk superadmin/HR manager).",
    "Atur konfigurasi umum aplikasi sesuai kebutuhan.",
    "Simpan perubahan setelah selesai melakukan pengaturan.",
])

# ---------------------------------------------------------------------------
# Struktur Proyek
# ---------------------------------------------------------------------------
doc.add_heading("2. Struktur Proyek", level=2)
para("Aplikasi Uranop menggunakan arsitektur monolith dengan pemisahan backend (Laravel) dan "
     "frontend (Vue 3). Backend disusun per modul bisnis, sedangkan frontend disusun per halaman "
     "(Pages) di bawah dua area akses, yaitu Admin dan Supervisor.")
code_block("""
app/Modules/
  Auth/          Autentikasi, user, role & permission, notifikasi
  Employee/      Data karyawan, kontrak, gaji, kompensasi, keluarga, dokumen, grouping
  Attendance/    Kehadiran, raw log, rekap, lembur, consecutive day
  Schedule/      Roster, pola kerja, shift, kalender
  Leave/         Cuti (generate, pengajuan, saldo, rekap)
  Kasbon/        Pengajuan, persetujuan, pelunasan kasbon
  Payroll/       Periode, gaji karyawan, slip gaji, THR, BPJS, PPh 21
  Reports/       Laporan dan ekspor
  Organization/  Bagian (departemen), jabatan, grade gaji
  Settings/      Pengaturan aplikasi
  Dashboard/     Ringkasan dashboard

resources/js/
  Pages/Admin/         Halaman area Admin
  Pages/Supervisor/    Halaman area Supervisor
  Pages/Auth/          Halaman login
  Components/          Komponen UI yang dipakai bersama
  Layouts/             Layout (Admin, Supervisor, Root)
  Stores/              State management (Pinia)
  router/index.js      Definisi rute & guard akses
""", lang_label="Struktur folder utama")

# ===========================================================================
# BAB IV PENUTUP
# ===========================================================================
page_break()
doc.add_heading("BAB IV PENUTUP", level=1)
doc.add_heading("Kesimpulan", level=2)
para("Berdasarkan pembahasan pada bab-bab sebelumnya, dapat ditarik kesimpulan bahwa aplikasi "
     "Uranop (Enterprise HRIS System) dibangun untuk menjawab kebutuhan pengelolaan sumber daya "
     "manusia yang terintegrasi dalam satu sistem berbasis web. Dengan adanya aplikasi ini, "
     "pengelolaan data karyawan, kehadiran, jadwal kerja, cuti, kasbon, BPJS, PPh 21, payroll, "
     "hingga pembuatan laporan dapat dilakukan secara lebih cepat, akurat, dan terpusat pada satu "
     "basis data.")
para("Aplikasi ini diharapkan dapat membantu perusahaan dalam mengambil keputusan yang lebih baik, "
     "meningkatkan kepatuhan terhadap ketentuan ketenagakerjaan dan perpajakan, serta mengurangi "
     "kesalahan manusia (human error) pada proses-proses yang sebelumnya dilakukan secara manual. "
     "Pengembangan lebih lanjut tetap terbuka, misalnya penambahan fitur analitik SDM, integrasi "
     "dengan perangkat absensi tambahan, serta aplikasi mobile bagi karyawan untuk pengajuan cuti "
     "dan kasbon secara mandiri.")

# ---------------------------------------------------------------------------
# Simpan
# ---------------------------------------------------------------------------
OUT = r"H:\laragon\www\instance1.uranop.com\Manual_Book_dan_Source_Code_Uranop.docx"
doc.save(OUT)
print("Berhasil membuat:", OUT)
